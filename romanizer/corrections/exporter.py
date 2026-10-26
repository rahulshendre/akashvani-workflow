"""Growing the dictionary: exporting approved engine finds for the experts.

This closes the loop. A correction a reviewer ticks today becomes an exact
dictionary match tomorrow, so the engine's guesswork shrinks over time.
"""
from __future__ import annotations

import io

from docx import Document


class DictionaryExporter:
    """Writes approved engine pairs in the dictionary's own table format."""

    def build(self, changes):
        """Make a 3-column docx (same format) of engine-found pairs to paste back.

        Exact dictionary hits are skipped — they are already in the dictionary.
        Returns (docx_stream, n_rows).
        """
        seen, rows = set(), []
        for ch in changes:
            if ch["tier"] == "exact":
                continue
            key = (ch["orig"].lower(), ch["repl"])
            if key in seen:
                continue
            seen.add(key)
            rows.append((ch["orig"], ch["repl"]))

        doc = Document()
        doc.add_heading("Engine-found corrections (review before merging)", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "", "AI Version", "Planet Read Version"
        for i, (a, b) in enumerate(rows, 1):
            cells = table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = str(i), a, b
        out = io.BytesIO()
        doc.save(out)
        out.seek(0)
        return out, len(rows)
