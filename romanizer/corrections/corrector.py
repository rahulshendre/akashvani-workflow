"""Scanning a romanized subtitle for corrections.

Produces CANDIDATES, not decisions. Exact dictionary matches are certain; the
engine's suggestions are roughly 75-80% precise on words outside the dictionary,
so they are returned tiered and flagged for a human rather than silently applied.
"""
from __future__ import annotations

import io
import logging
import re
import time

from docx import Document
from docx.oxml.ns import qn

from ..config import FIX_CUTOFF, SUGGEST_FLOOR, TIER_HL
from ..core.dictionary import Dictionary
from .engine import SuggestEngine
from .runs import RunBuilder

log = logging.getLogger("romanize.corrector")


class DocumentCorrector:
    """Finds every correction candidate in a romanized subtitle .docx."""

    WORD_RE = re.compile(r'[A-Za-z]+')

    def __init__(self, dictionary: Dictionary, enable_engine: bool = True):
        self.dictionary = dictionary
        self.enable_engine = enable_engine
        self.engine = SuggestEngine(dictionary.pairs) if enable_engine else None

    # ── one paragraph ────────────────────────────────────────────────────────
    def correct_paragraph(self, para, protected=frozenset()):
        """Return list of change dicts: {orig, repl, tier, conf, reason}.

        `protected` holds lowercase words that were pre-substituted from the
        dictionary's native column — they are already the confirmed PlanetRead
        spelling, so neither the exact dictionary nor the engine may touch them.
        """
        full_text = "".join(run.text for run in para.runs)
        if not full_text.strip():
            return []

        exact_re = self.dictionary.exact_regex
        replacements = self.dictionary.replacements

        # 1) exact dictionary spans — single regex pass, matches are non-overlapping
        spans = []  # (start, end, replacement, tier, conf, reason)
        if exact_re is not None:
            for m in exact_re.finditer(full_text):
                if m.group(1).lower() in protected:
                    continue
                rep = replacements.get(m.group(1).lower())
                if rep is not None:
                    spans.append((m.start(), m.end(), rep, "exact", 1.0, "dictionary"))

        # 2) engine suggestions on remaining words
        if self.enable_engine and self.engine is not None:
            for m in self.WORD_RE.finditer(full_text):
                if any(s < m.end() and m.start() < e for s, e, *_ in spans):
                    continue
                if m.group(0).lower() in protected:
                    continue
                sug = self.engine.suggest(m.group(0))
                if sug is None or sug.confidence < SUGGEST_FLOOR:
                    continue
                tier = "fix" if sug.confidence >= FIX_CUTOFF else "suggest"
                spans.append((m.start(), m.end(), sug.correction, tier,
                              sug.confidence, sug.reason))

        if not spans:
            return []

        spans.sort(key=lambda x: x[0])
        changes = [
            {"orig": full_text[s:e], "repl": rep, "tier": tier, "conf": conf, "reason": reason}
            for s, e, rep, tier, conf, reason in spans
        ]

        # rebuild the paragraph as: untouched | red original | highlighted fix | …
        first_run_xml = para.runs[0]._r if para.runs else None
        p_xml = para._p
        for r in p_xml.findall(qn('w:r')):
            p_xml.remove(r)

        pos = 0
        for (start, end, replacement, tier, conf, reason) in spans:
            if pos < start:
                p_xml.append(RunBuilder.normal(full_text[pos:start], first_run_xml))
            p_xml.append(RunBuilder.red(full_text[start:end], first_run_xml))
            p_xml.append(RunBuilder.normal(' ', first_run_xml))
            p_xml.append(RunBuilder.highlight(replacement, TIER_HL[tier], first_run_xml))
            pos = end
        if pos < len(full_text):
            p_xml.append(RunBuilder.normal(full_text[pos:], first_run_xml))

        return changes

    # ── whole document ───────────────────────────────────────────────────────
    def run(self, input_bytes: bytes, progress_cb=None, protected=frozenset()):
        """Scan the romanized subtitle .docx.

        `progress_cb(done, total, stage_msg)` is called as work proceeds (optional).
        `protected`: lowercase pre-substituted words that must not be changed.
        Returns (marked_docx_stream, replacements, all_changes).
        """
        t_start = time.perf_counter()

        def report(done, total, msg):
            if progress_cb:
                progress_cb(done, total, msg)

        report(0, 1, "Reading dictionary…")
        replacements = self.dictionary.replacements

        report(0, 1, "Training pattern engine…")
        # engine was built in __init__; this stage message is kept for the UI

        report(0, 1, "Opening subtitle file…")
        doc = Document(io.BytesIO(input_bytes))

        # gather every paragraph up-front so we know the total for the progress bar
        paras = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paras.extend(cell.paragraphs)
        total = len(paras)
        log.info("Scanning %d paragraphs (engine=%s)…", total,
                 "on" if self.enable_engine else "off")

        all_changes = []
        t_scan = time.perf_counter()
        for i, para in enumerate(paras, 1):
            all_changes.extend(self.correct_paragraph(para, protected))
            if i % 200 == 0 or i == total:
                log.info("  …%d/%d paragraphs, %d changes so far", i, total, len(all_changes))
                report(i, total, f"Scanning paragraphs… {i}/{total}")
        log.info("Scan done: %d changes in %.2fs", len(all_changes),
                 time.perf_counter() - t_scan)

        if self.engine is not None:
            log.info("Engine cache: %d unique words looked up", len(self.engine._cache))

        report(total, total, "Saving fixed file…")
        out = io.BytesIO()
        doc.save(out)
        out.seek(0)
        log.info("DocumentCorrector finished in %.2fs total", time.perf_counter() - t_start)
        return out, replacements, all_changes
